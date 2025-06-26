import json
from abc import ABC, abstractmethod

import openpyxl


class WorkWithFile(ABC):
    """
    Абстрактный класс для работы с файлом.

    Класс содержит методы, которые будут реализованы в классах наследниках.

    А именно:
    Абстрактный метод get_data, выполняющий получение данных из файла.
    Абстрактный метод addition_data, выполняющий добавление данных в файл.
    Абстрактный метод del_data, выполняющий удаление данных из файла.
    """

    @abstractmethod
    def get_data(self):
        """Метод получения данных из файла"""
        pass

    @abstractmethod
    def addition_data(self, vacancy_dict):
        """Метод добавления данных в файл"""
        pass

    @abstractmethod
    def del_data(self, criterion_key, criterion_value):
        """Метод удаления данных из файла"""
        pass


class work_with_json(WorkWithFile):
    """
    Класс для работы с JSON-файлами.

    Класс содержит следующие методы:

    Метод get_data, выполняющий получение данных из файла.
    Метод addition_data, выполняет добавление новой вакансии без дублирования.
    Метод del_data, выполняет удаление вакансии по критерию.
    """

    def __init__(self, filename="vacancies.json"):
        self.__filename = filename  # приватный атрибут

    def get_data(self) -> list[dict]:
        """Получение данных из файла"""
        try:
            with open(self.__filename, "r", encoding="utf-8") as f:
                data = json.load(f)
                return data  # список словарей
        except (FileNotFoundError, json.JSONDecodeError):
            return []

    def addition_data(self, vacancy_dict: dict):
        """
        Метод добавления новой вакансии без дублирования.
        Проверка по ключу 'name' или другому уникальному ключу.
        """
        data = self.get_data()
        # Предположим, что уникальный ключ — 'name'
        vacancy_name = vacancy_dict.get("name")
        if not any(vac.get("name") == vacancy_name for vac in data):
            data.append(vacancy_dict)
            with open(self.__filename, "w", encoding="utf-8") as f:
                json.dump(data, f, indent=4, ensure_ascii=False)

    def del_data(self, criterion_key: str, criterion_value: str):
        """
        Удаление вакансии по критерию.
        Например: criterion_key='name', criterion_value='Python Developer'
        """
        data = self.get_data()
        new_data = [vac for vac in data if vac.get(criterion_key) != criterion_value]
        with open(self.__filename, "w", encoding="utf-8") as f:
            json.dump(new_data, f, indent=4, ensure_ascii=False)


# Дополнительные классы для работы с файлом.


class work_with_excel(WorkWithFile):
    """
    Класс для работы с Excel-файлами

    Класс содержит следующие методы:

    Метод get_data, выполняющий получение данных из файла.
    Метод addition_data, выполняет добавление новой вакансии без дублирования.
    Метод del_data, выполняет удаление вакансии по критерию.
    """

    def __init__(self, filename="vacancies.xlsx"):
        self.__filename = filename

    def get_data(self) -> list[dict]:
        """Получение данных из Excel файла"""
        try:
            wb = openpyxl.load_workbook(self.__filename)
            sheet = wb.active
            data = []
            headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
            for row in sheet.iter_rows(min_row=2, values_only=True):
                row_dict = {headers[i]: row[i] for i in range(len(headers))}
                data.append(row_dict)
            return data
        except FileNotFoundError:
            return []

    def addition_data(self, vacancy_dict: dict):
        """
        Добавление новой вакансии без дублирования по ключу 'name'.
        """
        data = self.get_data()
        vacancy_name = vacancy_dict.get("name")
        if not any(vac.get("name") == vacancy_name for vac in data):
            try:
                wb = openpyxl.load_workbook(self.__filename)
                sheet = wb.active
            except FileNotFoundError:
                wb = openpyxl.Workbook()
                sheet = wb.active
                headers = list(vacancy_dict.keys())
                sheet.append(headers)
            else:
                headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]

            row_values = [vacancy_dict.get(header) for header in headers]
            sheet.append(row_values)
            wb.save(self.__filename)

    def del_data(self, criterion_key: str, criterion_value: str):
        """
        Удаление вакансий по критерию.
        """
        try:
            wb = openpyxl.load_workbook(self.__filename)
            sheet = wb.active
            headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
            key_index = headers.index(criterion_key)

            rows_to_delete = []
            for idx, row in enumerate(sheet.iter_rows(min_row=2), start=2):
                cell_value = row[key_index].value
                if cell_value == criterion_value:
                    rows_to_delete.append(idx)

            for row_idx in reversed(rows_to_delete):
                sheet.delete_rows(row_idx)

            wb.save(self.__filename)
        except FileNotFoundError:
            pass


class WorkWithTXT(WorkWithFile):
    """
    Класс для работы с TXT-файлами

    Класс содержит следующие методы:

    Метод get_data, выполняющий получение данных из файла.
    Метод addition_data, выполняет добавление новой вакансии без дублирования.
    Метод del_data, выполняет удаление вакансии по критерию.
    """

    def __init__(self, filename="vacancies.txt"):
        self.__filename = filename

    def get_data(self) -> list[dict]:
        """Получение данных из TXT файла"""
        data = []
        try:
            with open(self.__filename, "r", encoding="utf-8") as f:
                for line in f:
                    line = line.strip()
                    if line:
                        vacancy_dict = {}
                        parts = line.split(";")
                        for part in parts:
                            if ":" in part:
                                key, value = part.split(":", 1)
                                vacancy_dict[key.strip()] = value.strip()
                        if vacancy_dict:
                            data.append(vacancy_dict)
        except FileNotFoundError:
            return []
        return data

    def addition_data(self, vacancy_dict: dict):
        """
        Добавление новой вакансии.
        Формат строки: ключ: значение; ключ: значение; ...
        """
        # Формируем строку из словаря
        line_parts = [f"{key}: {value}" for key, value in vacancy_dict.items()]
        line_str = "; ".join(line_parts)

        data = self.get_data()
        vacancy_name = vacancy_dict.get("name")

        if not any(vac.get("name") == vacancy_name for vac in data):
            with open(self.__filename, "a", encoding="utf-8") as f:
                f.write(line_str + "\n")

    def del_data(self, criterion_key: str, criterion_value: str):
        """Удаление вакансий по критерию"""
        data = self.get_data()

        new_data = [vac for vac in data if vac.get(criterion_key) != criterion_value]

        with open(self.__filename, "w", encoding="utf-8") as f:
            for vac in new_data:
                line_parts = [f"{key}: {value}" for key, value in vac.items()]
                line_str = "; ".join(line_parts)
                f.write(line_str + "\n")
